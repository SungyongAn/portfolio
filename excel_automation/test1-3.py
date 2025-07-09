from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

wb = load_workbook('/Users/sungyongan/Desktop/sample2.xlsx')
ws = wb.active

# テーブルを生成する
table = Table(displayName='Table1', ref='A1:D500')

# テーブルのスタイルを決めておく
table_style = TableStyleInfo(name='TableStyleMedium8', showRowStripes=True)

# テーブルのスタイルを設定
table.tableStyleInfo = table_style

# シートにテーブルを追加
ws.add_table(table)

wb.save('/Users/sungyongan/Desktop/sample2_テーブル.xlsx')
